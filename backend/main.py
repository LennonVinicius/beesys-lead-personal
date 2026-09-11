import os
import threading
import time
import base64
import io
from datetime import datetime, timedelta, timezone
from typing import Literal, Optional

from fastapi import Depends, FastAPI, HTTPException, Query, Response
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field

from auth_api import get_current_user
from db import (
    add_activity, ai_usage_today, cancel_search_job, complete_followup, create_followup, create_goal,
    create_search_job, dashboard_advanced, dashboard_summary, get_app_setting, get_business, get_search_job,
    goal_progress, init_db, list_businesses, list_businesses_by_city, list_businesses_by_district,
    list_campaigns, list_search_jobs, list_team_members, load_activities, load_followups, load_goals, latest_search_job_for_business,
    load_job_businesses, load_objections, quick_outcome, record_objection, set_app_setting, update_crm, update_visit,
    upsert_businesses, recalculate_business_score, resolve_identity_rows, load_lead_changes, duplicate_contact_candidates,
    list_cadence_templates, save_cadence_template, assign_cadence, list_lead_cadences,
    list_automation_rules, save_automation_rule, set_automation_active, list_webhooks, save_webhook,
    list_user_roles, upsert_user_role, record_import_batch, list_import_batches, create_privacy_request, list_privacy_requests, complete_privacy_request, list_recent_lead_changes,
)
from jobs import process_search_job_batch
from providers import (
    ProviderTemporaryError, geocode_geoapify, merge_provider_rows, search_foursquare, search_geoapify, search_osm,
)
from reporting import build_business_report
from routing import plan_timed_route, plan_route_advanced
from targeting import annotate_target_fit
from analyzer import analyze_site
from operations import (
    audit, list_audit_logs, enrich_lead, today_workspace, create_public_report, get_public_report,
    build_post_visit_message, business_evolution, schedule_reanalysis, due_reanalysis, mark_reanalyzed,
    create_route_record, get_route, list_routes, update_route_stops, update_route_stop_status,
    create_territory, list_territories, upsert_icp, list_icps, export_full, campaign_financials, record_provider_usage, provider_usage_summary,
)
from intelligence import intelligence_overview, natural_filter, copilot_answer, compute_lead_intelligence, competitor_playbook
from identity import identity_key as compute_identity_key
from webhooks import validate_webhook_url, dispatch_event, list_deliveries
from automation_engine import run_automations

app = FastAPI(title="BeeSys Lead Search API", version="2.0.0")
_ACTIVE_JOBS = set()
_ACTIVE_LOCK = threading.Lock()

origins = [x.strip().rstrip("/") for x in (os.getenv("FRONTEND_ORIGINS") or "http://localhost:5173").split(",") if x.strip()]
# Produção fixa + previews do projeto na Vercel. O regex é restrito ao nome do app.
app.add_middleware(
    CORSMiddleware,
    allow_origins=origins,
    allow_origin_regex=os.getenv("FRONTEND_ORIGIN_REGEX", r"https://beesys-lead-[a-zA-Z0-9-]+\.vercel\.app"),
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

GEOAPIFY_KEY = os.getenv("GEOAPIFY_API_KEY") or ""
FOURSQUARE_KEY = os.getenv("FOURSQUARE_API_KEY") or ""
_AI_KEYS = {
    "Gemini": os.getenv("GEMINI_API_KEY") or "",
    "OpenRouter": os.getenv("OPENROUTER_API_KEY") or "",
    "OpenAI": os.getenv("OPENAI_API_KEY") or "",
}
DEFAULT_AI_MODELS = {
    "Gemini": os.getenv("GEMINI_MODEL", "gemini-3.7-flash"),
    "OpenRouter": os.getenv("OPENROUTER_MODEL", "openrouter/free"),
    "OpenAI": os.getenv("OPENAI_MODEL", "gpt-5.6-luna"),
}
DEFAULT_SETTINGS = {
    "default_mrr": 89.90,
    "default_score_profile": "Auto",
    "default_ai_provider": "Desativada",
    "ai_min_score": int(os.getenv("AI_MIN_SCORE", "60")),
    "ai_daily_call_limit": int(os.getenv("AI_DAILY_CALL_LIMIT", "100")),
    "ai_daily_budget_usd": float(os.getenv("AI_DAILY_BUDGET_USD", "2")),
    "ai_limit_per_batch": int(os.getenv("AI_LIMIT_PER_BATCH", "8")),
    "gemini_model": DEFAULT_AI_MODELS["Gemini"],
    "openrouter_model": DEFAULT_AI_MODELS["OpenRouter"],
    "openai_model": DEFAULT_AI_MODELS["OpenAI"],
    "route_mode": "driving-car",
    "route_available_minutes": 240,
    "route_visit_minutes": 15,
    "route_max_stops": 12,
}


class SearchIn(BaseModel):
    query_text: str = "Área selecionada no mapa"
    lat: float
    lon: float
    radius_m: int = Field(1000, ge=100, le=15000)
    provider: Literal["auto", "geoapify", "osm", "foursquare"] = "auto"
    max_results: int = Field(300, ge=10, le=1000)
    auto_analyze: bool = True
    ai_provider: Literal["Desativada", "Auto", "Gemini", "OpenRouter", "OpenAI"] = "Desativada"
    ai_model: Optional[str] = None
    score_profile: str = "Auto"
    campaign_name: str = "Prospecção geral"
    default_mrr: float = 89.90
    city_name: Optional[str] = None
    state_name: Optional[str] = None
    district_name: Optional[str] = None
    include_large_chains: bool = False
    icp_name: Optional[str] = None
    territory_id: Optional[int] = None
    incremental_from_job_id: Optional[int] = None


class CRMIn(BaseModel):
    pipeline_status: str = "NEW"
    assigned_to: str = ""
    contact_name: str = ""
    contact_phone: str = ""
    next_action_at: Optional[str] = None
    notes: str = ""
    lost_reason: str = ""
    do_not_contact: bool = False
    estimated_mrr: float = 89.90


class ActivityIn(BaseModel):
    activity_type: str
    details: str = ""
    outcome: str = ""
    lat: Optional[float] = None
    lon: Optional[float] = None


class ObjectionIn(BaseModel):
    code: str
    label: str
    details: str = ""


class OutcomeIn(BaseModel):
    outcome: Literal["VISITED", "NO_OWNER", "INTERESTED", "DEMO", "PROPOSAL", "CLIENT", "NOT_INTERESTED"]
    notes: str = ""
    objection_code: Optional[str] = None
    objection_label: Optional[str] = None
    objection_details: str = ""


class FollowupIn(BaseModel):
    business_key: str
    title: str
    due_at: str
    notes: str = ""


class GoalIn(BaseModel):
    metric: str
    target: float
    start_at: str
    end_at: str
    period: str = "DAILY"
    label: str = ""


class RouteIn(BaseModel):
    origin_lat: float
    origin_lon: float
    origin_label: str = "Minha localização"
    end_lat: Optional[float] = None
    end_lon: Optional[float] = None
    end_label: Optional[str] = None
    business_keys: list[str]
    fixed_business_keys: list[str] = Field(default_factory=list)
    excluded_business_keys: list[str] = Field(default_factory=list)
    mode: Literal["driving-car", "foot-walking"] = "driving-car"
    strategy: Literal["balanced", "sales", "visits", "distance", "manual"] = "balanced"
    available_minutes: int = Field(240, ge=30, le=720)
    visit_minutes: int = Field(15, ge=5, le=120)
    max_stops: int = Field(12, ge=1, le=50)
    return_to_start: bool = False
    persist: bool = False
    name: Optional[str] = None
    route_date: Optional[str] = None
    campaign_name: Optional[str] = None
    start_local: Optional[str] = None

class RouteStopsIn(BaseModel):
    stops: list[dict]

class RouteStopStatusIn(BaseModel):
    status: str
    reason: str = ""

class PublicReportIn(BaseModel):
    job_id: Optional[int] = None
    expires_days: int = Field(30, ge=1, le=365)

class TerritoryIn(BaseModel):
    name: str
    city_name: Optional[str] = None
    district_name: Optional[str] = None
    assigned_to: Optional[str] = None
    center_lat: Optional[float] = None
    center_lon: Optional[float] = None
    radius_m: Optional[int] = None
    polygon: Optional[list] = None

class ICPIn(BaseModel):
    name: str
    config: dict = {}

class ReanalysisIn(BaseModel):
    days: int = Field(30, ge=1, le=365)

class MessageIn(BaseModel):
    notes: str = ""


class CopilotIn(BaseModel):
    question: str = Field(..., min_length=2, max_length=1000)

class NaturalQueryIn(BaseModel):
    query: str = Field(..., min_length=2, max_length=1000)

class CadenceIn(BaseModel):
    name: str
    description: str = ""
    steps: list[dict] = Field(default_factory=list)

class AssignCadenceIn(BaseModel):
    cadence_id: int

class AutomationIn(BaseModel):
    name: str
    event_name: str = "lead.updated"
    conditions: dict = Field(default_factory=dict)
    actions: list[dict] = Field(default_factory=list)
    active: bool = True

class AutomationToggleIn(BaseModel):
    active: bool

class WebhookIn(BaseModel):
    name: str
    url: str
    events: list[str] = Field(default_factory=lambda:["lead.client"])
    secret: Optional[str] = None

class RoleIn(BaseModel):
    email: str
    role: Literal["ADMIN","MANAGER","SELLER","VIEWER"] = "SELLER"

class ImportLeadsIn(BaseModel):
    filename: str = "import.csv"
    rows: list[dict] = Field(default_factory=list, max_length=5000)

class ImportFileIn(BaseModel):
    filename: str
    content_base64: str = Field(..., min_length=4)

class PrivacyIn(BaseModel):
    request_type: str = "REVIEW"
    business_key: Optional[str] = None
    contact_value: Optional[str] = None
    details: str = ""

class PrivacyCompleteIn(BaseModel):
    anonymize: bool = False

class SettingsIn(BaseModel):
    default_mrr: float = 89.90
    default_score_profile: str = "Auto"
    default_ai_provider: str = "Desativada"
    ai_min_score: int = 60
    ai_daily_call_limit: int = 100
    ai_daily_budget_usd: float = 2.0
    ai_limit_per_batch: int = 8
    gemini_model: str = DEFAULT_AI_MODELS["Gemini"]
    openrouter_model: str = DEFAULT_AI_MODELS["OpenRouter"]
    openai_model: str = DEFAULT_AI_MODELS["OpenAI"]
    route_mode: str = "driving-car"
    route_available_minutes: int = 240
    route_visit_minutes: int = 15
    route_max_stops: int = 12


def _user_email(user):
    return (user.get("email") or "").lower()

def _role_for_user(user):
    email=_user_email(user)
    roles=list_user_roles()
    if not roles:
        return 'ADMIN'
    match=next((r for r in roles if (r.get('email') or '').lower()==email),None)
    return (match.get('role') if match else 'SELLER') or 'SELLER'


def _require_role(user, allowed):
    role=_role_for_user(user)
    if role not in set(allowed):
        raise HTTPException(403, 'Seu papel não permite esta ação')
    return role


def _settings():
    saved = get_app_setting("lead_search_settings", {}) or {}
    return {**DEFAULT_SETTINGS, **saved}


def _ai_runtime():
    settings = _settings()
    models = {
        "Gemini": settings.get("gemini_model") or DEFAULT_AI_MODELS["Gemini"],
        "OpenRouter": settings.get("openrouter_model") or DEFAULT_AI_MODELS["OpenRouter"],
        "OpenAI": settings.get("openai_model") or DEFAULT_AI_MODELS["OpenAI"],
    }
    policy = {
        "min_score": int(settings.get("ai_min_score") or 0),
        "daily_call_limit": int(settings.get("ai_daily_call_limit") or 0),
        "daily_budget_usd": float(settings.get("ai_daily_budget_usd") or 0),
    }
    return models, policy, settings


def _apply_icp(rows, icp_name):
    if not icp_name:
        return rows
    profile = next((x for x in list_icps() if x.get("name") == icp_name), None)
    if not profile:
        return rows
    cfg = profile.get("config") or {}
    include_terms=[str(x).lower() for x in cfg.get("include_terms",[]) if x]
    exclude_terms=[str(x).lower() for x in cfg.get("exclude_terms",[]) if x]
    min_fit=int(cfg.get("min_target_fit") or 0)
    out=[]
    for r in rows:
        blob=f"{r.get('name','')} {r.get('category','')}".lower()
        if include_terms and not any(t in blob for t in include_terms):
            continue
        if exclude_terms and any(t in blob for t in exclude_terms):
            continue
        if int(r.get("target_fit_score") or 0) < min_fit:
            continue
        if cfg.get("local_only") and r.get("is_large_chain"):
            continue
        r["icp_name"] = icp_name
        out.append(r)
    return out


def _discover(payload: SearchIn):
    groups = []
    source_counts = {}
    raw_count = 0
    if payload.provider in {"auto", "geoapify"} and GEOAPIFY_KEY:
        geo_rows=search_geoapify(payload.lat, payload.lon, payload.radius_m, GEOAPIFY_KEY, payload.max_results)
        groups.append(geo_rows); source_counts["Geoapify"]=len(geo_rows); raw_count += len(geo_rows)
        record_provider_usage("Geoapify","places",1,len(geo_rows),{"radius_m":payload.radius_m})
    if payload.provider == "osm" or (payload.provider == "auto" and (not groups or len(groups[0]) < 30)):
        try:
            osm_rows=search_osm(payload.lat, payload.lon, payload.radius_m, payload.max_results)
            groups.append(osm_rows); source_counts["OpenStreetMap"]=len(osm_rows); raw_count += len(osm_rows)
            record_provider_usage("OpenStreetMap","overpass",1,len(osm_rows),{"radius_m":payload.radius_m})
        except ProviderTemporaryError:
            if not groups:
                raise
    if payload.provider == "foursquare":
        if not FOURSQUARE_KEY:
            raise HTTPException(400, "FOURSQUARE_API_KEY não configurada")
        fsq_rows=search_foursquare(payload.lat, payload.lon, payload.radius_m, FOURSQUARE_KEY, max_results=payload.max_results)
        groups.append(fsq_rows); source_counts["Foursquare"]=len(fsq_rows); raw_count += len(fsq_rows)
        record_provider_usage("Foursquare","places",1,len(fsq_rows),{"radius_m":payload.radius_m})
    if payload.provider == "geoapify" and not GEOAPIFY_KEY:
        raise HTTPException(400, "GEOAPIFY_API_KEY não configurada")

    rows = merge_provider_rows(*groups)[:payload.max_results]
    resolve_identity_rows(rows)
    dedup_count=len(rows)
    incremental_filtered=0
    if payload.incremental_from_job_id:
        previous=get_search_job(payload.incremental_from_job_id)
        if previous:
            from providers import haversine_m
            prev_radius=int(previous.get('radius_m') or 0)
            same_center=haversine_m(payload.lat,payload.lon,float(previous.get('center_lat') or payload.lat),float(previous.get('center_lon') or payload.lon)) <= 100
            if same_center and payload.radius_m > prev_radius:
                before=len(rows)
                rows=[r for r in rows if haversine_m(payload.lat,payload.lon,float(r['lat']),float(r['lon'])) > prev_radius*0.97]
                incremental_filtered=before-len(rows)
    annotate_target_fit(rows)
    prepared = []
    for r in rows:
        r["campaign_name"] = payload.campaign_name.strip() or "Prospecção geral"
        r["estimated_mrr"] = payload.default_mrr
        r["city_name"] = r.get("city_name") or payload.city_name
        r["state_name"] = r.get("state_name") or payload.state_name
        r["district_name"] = r.get("district_name") or payload.district_name
        if not payload.include_large_chains and r.get("is_large_chain"):
            continue
        prepared.append(r)
    prepared=_apply_icp(prepared, payload.icp_name)
    meta={"source_counts":source_counts,"raw_discovered_count":raw_count,"deduplicated_count":dedup_count,"incremental_filtered":incremental_filtered}
    return prepared,meta


def _drain_job(job_id: int):
    with _ACTIVE_LOCK:
        if job_id in _ACTIVE_JOBS:
            return
        _ACTIVE_JOBS.add(job_id)
    try:
        for _ in range(1000):
            job = get_search_job(job_id)
            if not job or job.get("status") in {"DONE", "FAILED", "CANCELED"}:
                return
            models, policy, _ = _ai_runtime()
            cfg = job.get("config") or {}
            selected_provider = cfg.get("ai_provider")
            if cfg.get("ai_model") and selected_provider in models:
                models[selected_provider] = cfg["ai_model"]
            process_search_job_batch(
                job_id,
                batch_size=int(os.getenv("JOB_BATCH_SIZE", "20")),
                ai_keys=_AI_KEYS,
                ai_models=models,
                ai_policy=policy,
            )
            time.sleep(0.05)
    finally:
        with _ACTIVE_LOCK:
            _ACTIVE_JOBS.discard(job_id)


def _recovery_loop():
    while True:
        try:
            for job in list_search_jobs(50):
                if job.get("status") in {"RUNNING", "QUEUED"}:
                    threading.Thread(target=_drain_job, args=(int(job["id"]),), daemon=True).start()
        except Exception:
            pass
        time.sleep(8)


def _reanalysis_loop():
    # Pequenos lotes para não competir com as buscas ativas no Render Free.
    while True:
        try:
            for lead in due_reanalysis(3):
                if not lead.get("website"):
                    continue
                try:
                    result = analyze_site(lead["website"], use_cache=False)
                    lead.update(result)
                    upsert_businesses([lead])
                    recalculate_business_score(lead["business_key"])
                    mark_reanalyzed(lead["business_key"], 60)
                    add_activity(lead["business_key"], "REANALYSIS", "Reanálise automática da presença digital", "DONE", "system")
                except Exception:
                    # Agenda para uma nova tentativa em 7 dias, sem travar o backend.
                    schedule_reanalysis(lead["business_key"], 7)
        except Exception:
            pass
        time.sleep(int(os.getenv("REANALYSIS_SCAN_SECONDS", "1800")))


@app.on_event("startup")
def startup():
    init_db()
    threading.Thread(target=_recovery_loop, daemon=True).start()
    threading.Thread(target=_reanalysis_loop, daemon=True).start()


@app.get("/health")
def health():
    return {"status": "ok", "service": "beesys-lead-search"}


@app.get("/api/geocode")
def geocode(q: str = Query(min_length=2), user=Depends(get_current_user)):
    result = geocode_geoapify(q, GEOAPIFY_KEY)
    if not result:
        raise HTTPException(404, "Local não encontrado")
    return result


@app.get('/api/me/role')
def my_role(user=Depends(get_current_user)):
    return {'email':_user_email(user),'role':_role_for_user(user)}


@app.get("/api/settings")
def settings_get(user=Depends(get_current_user)):
    s = _settings()
    return {**s, "ai_usage_today": ai_usage_today(), "available_ai": {k: bool(v) for k, v in _AI_KEYS.items()}}


@app.patch("/api/settings")
def settings_update(payload: SettingsIn, user=Depends(get_current_user)):
    _require_role(user, {'ADMIN','MANAGER'})
    data = payload.model_dump()
    set_app_setting("lead_search_settings", data)
    audit(_user_email(user), "UPDATE_SETTINGS", "system", "lead_search_settings", data)
    return {**data, "ai_usage_today": ai_usage_today(), "available_ai": {k: bool(v) for k, v in _AI_KEYS.items()}}


@app.get("/api/campaigns")
def campaigns(user=Depends(get_current_user)):
    return list_campaigns()


@app.get("/api/team")
def team(user=Depends(get_current_user)):
    members = list_team_members()
    email = _user_email(user)
    if email and email not in members:
        members.insert(0, email)
    return members


@app.post("/api/searches")
def create_search(payload: SearchIn, user=Depends(get_current_user)):
    try:
        rows, discovery_meta = _discover(payload)
    except ProviderTemporaryError as exc:
        raise HTTPException(503, str(exc))
    _, _, settings = _ai_runtime()
    center = {"lat": payload.lat, "lon": payload.lon, "display_name": payload.query_text}
    config = {
        "auto_analyze": payload.auto_analyze,
        "ai_provider": payload.ai_provider,
        "ai_model": payload.ai_model,
        "score_profile": payload.score_profile,
        "cache_days": 14,
        "ai_limit_per_batch": int(settings.get("ai_limit_per_batch") or 8),
        "campaign_name": payload.campaign_name,
        "city_name": payload.city_name,
        "state_name": payload.state_name,
        "district_name": payload.district_name,
        "default_mrr": payload.default_mrr,
        "include_large_chains": payload.include_large_chains,
        "icp_name": payload.icp_name,
        "territory_id": payload.territory_id,
        "incremental_from_job_id": payload.incremental_from_job_id,
        "discovery_meta": discovery_meta,
    }
    job_id = create_search_job(payload.query_text, center, payload.radius_m, payload.provider, rows, config, _user_email(user))
    threading.Thread(target=_drain_job, args=(job_id,), daemon=True).start()
    return {"job_id": job_id, "discovered": len(rows), **discovery_meta}


@app.get("/api/searches")
def searches(limit: int = 30, user=Depends(get_current_user)):
    return list_search_jobs(limit)


@app.get("/api/searches/{job_id}")
def search_detail(job_id: int, user=Depends(get_current_user)):
    job = get_search_job(job_id)
    if not job:
        raise HTTPException(404, "Busca não encontrada")
    job["results"] = [enrich_lead(r) for r in load_job_businesses(job_id, limit=1000)]
    return job


@app.post("/api/searches/{job_id}/cancel")
def search_cancel(job_id: int, user=Depends(get_current_user)):
    cancel_search_job(job_id)
    return {"ok": True}


@app.get("/api/leads")
def leads(
    limit: int = 500,
    status: Optional[str] = None,
    min_score: Optional[int] = None,
    search: Optional[str] = None,
    campaign: Optional[str] = None,
    provider: Optional[str] = None,
    assigned_to: Optional[str] = None,
    category: Optional[str] = None,
    website: Optional[Literal["any", "yes", "no", "problem"]] = "any",
    booking: Optional[Literal["any", "yes", "no", "manual"]] = "any",
    visited: Optional[Literal["any", "yes", "no"]] = "any",
    include_large_chains: bool = False,
    user=Depends(get_current_user),
):
    rows = list_businesses(max(limit * 4, 1000), status, min_score, search)
    def keep(r):
        if campaign and r.get("campaign_name") != campaign: return False
        if provider and r.get("provider") != provider: return False
        if assigned_to and r.get("assigned_to") != assigned_to: return False
        if category and category.lower() not in str(r.get("category") or "").lower(): return False
        if not include_large_chains and r.get("is_large_chain"): return False
        if website == "yes" and not r.get("website"): return False
        if website == "no" and r.get("website"): return False
        if website == "problem" and not (r.get("website") and r.get("site_functional") is False): return False
        if booking == "yes" and not r.get("has_booking"): return False
        if booking == "no" and (r.get("has_booking") or r.get("manual_booking_detected")): return False
        if booking == "manual" and not r.get("manual_booking_detected"): return False
        if visited == "yes" and not r.get("visited"): return False
        if visited == "no" and r.get("visited"): return False
        return True
    return [enrich_lead(r) for r in rows if keep(r)][:limit]


@app.get("/api/leads/{business_key}")
def lead_detail(business_key: str, user=Depends(get_current_user)):
    lead = get_business(business_key)
    if not lead:
        raise HTTPException(404, "Estabelecimento não encontrado")
    lead["activities"] = load_activities(business_key, 200)
    lead["objections"] = load_objections(business_key, 100)
    lead["followups"] = [f for f in load_followups(None, 1000) if f.get("business_key") == business_key][:100]
    lead["changes"] = load_lead_changes(business_key, 100)
    lead["cadences"] = list_lead_cadences(business_key)
    lead["duplicates"] = duplicate_contact_candidates(business_key, lead.get("phone") or lead.get("contact_phone"), lead.get("website"), lead.get("name"), 8)
    lead["intelligence"] = compute_lead_intelligence(lead)
    lead["competitor_playbook"] = competitor_playbook(lead)
    return enrich_lead(lead)


@app.patch("/api/leads/{business_key}")
def lead_update(business_key: str, payload: CRMIn, user=Depends(get_current_user)):
    if not get_business(business_key):
        raise HTTPException(404, "Estabelecimento não encontrado")
    update_crm(
        business_key, pipeline_status=payload.pipeline_status, assigned_to=payload.assigned_to,
        contact_name=payload.contact_name, contact_phone=payload.contact_phone, next_action_at=payload.next_action_at,
        notes=payload.notes, do_not_contact=payload.do_not_contact, lost_reason=payload.lost_reason,
        estimated_mrr=payload.estimated_mrr, actor_email=_user_email(user),
    )
    add_activity(business_key, "CRM_UPDATE", payload.notes, payload.pipeline_status, _user_email(user))
    audit(_user_email(user), "UPDATE_CRM", "business", business_key, payload.model_dump())
    event="lead.client" if payload.pipeline_status=="CLIENT" else "lead.updated"
    threading.Thread(target=run_automations,args=(event,business_key,_user_email(user)),daemon=True).start()
    threading.Thread(target=dispatch_event,args=(event,{"business_key":business_key,"pipeline_status":payload.pipeline_status}),daemon=True).start()
    return enrich_lead(get_business(business_key))


@app.post("/api/leads/{business_key}/activities")
def activity(business_key: str, payload: ActivityIn, user=Depends(get_current_user)):
    add_activity(business_key, payload.activity_type, payload.details, payload.outcome, _user_email(user), payload.lat, payload.lon)
    audit(_user_email(user), "ADD_ACTIVITY", "business", business_key, payload.model_dump())
    return {"ok": True}


@app.post("/api/leads/{business_key}/objections")
def objection(business_key: str, payload: ObjectionIn, user=Depends(get_current_user)):
    record_objection(business_key, payload.code, payload.label, payload.details, _user_email(user))
    audit(_user_email(user), "ADD_OBJECTION", "business", business_key, payload.model_dump())
    return {"ok": True}


@app.post("/api/leads/{business_key}/outcomes")
def outcome(business_key: str, payload: OutcomeIn, user=Depends(get_current_user)):
    status = quick_outcome(business_key, payload.outcome, payload.notes, _user_email(user))
    if payload.objection_code and payload.objection_label:
        record_objection(business_key, payload.objection_code, payload.objection_label, payload.objection_details, _user_email(user))
    audit(_user_email(user), "QUICK_OUTCOME", "business", business_key, payload.model_dump())
    event="lead.client" if status=="CLIENT" else "lead.outcome"
    threading.Thread(target=run_automations,args=(event,business_key,_user_email(user)),daemon=True).start()
    threading.Thread(target=dispatch_event,args=(event,{"business_key":business_key,"pipeline_status":status,"outcome":payload.outcome}),daemon=True).start()
    return {"ok": True, "pipeline_status": status}


@app.post("/api/leads/{business_key}/visit")
def visit(business_key: str, notes: str = "", user=Depends(get_current_user)):
    update_visit(business_key, True, notes)
    add_activity(business_key, "VISIT", notes, "VISITED", _user_email(user))
    return {"ok": True}


@app.get("/api/reports/business/{business_key}")
def business_report(business_key: str, job_id: Optional[int] = None, user=Depends(get_current_user)):
    lead = get_business(business_key)
    if not lead:
        raise HTTPException(404, "Estabelecimento não encontrado")
    job = get_search_job(job_id) if job_id is not None else latest_search_job_for_business(business_key)
    cfg = (job or {}).get("config") or {}
    city = lead.get("city_name") or cfg.get("city_name")
    district = lead.get("district_name") or cfg.get("district_name")
    # Quando o lead não possui mais o job original (por exemplo, importado ou
    # migrado), ainda geramos um relatório útil com a base histórica da cidade/área.
    peers = load_job_businesses(job["id"], 5000) if job else list_businesses_by_city(city, 5000)
    if not peers:
        peers = [lead]
    report=build_business_report(
        enrich_lead(lead), peers, job,
        city_peers=list_businesses_by_city(city, 5000),
        district_peers=list_businesses_by_district(city, district, 5000),
    )
    report["evolution"]=business_evolution(business_key)
    report["solution_map"]={
        "NO_WEBSITE":"Página própria BeeSys com conteúdo rastreável e marca do negócio",
        "SITE_PROBLEM":"Reestruturação da presença própria e página funcional",
        "MANUAL_BOOKING":"Agenda BeeSys para disponibilidade, confirmações e menos troca manual de mensagens",
        "NO_BOOKING":"Agendamento online BeeSys integrado ao fluxo do estabelecimento",
        "NO_CATALOG":"Catálogo/serviços BeeSys para apresentar oferta antes do contato",
        "NO_STRUCTURED_DATA":"Melhoria de estrutura e dados do negócio na presença própria",
        "COMPETITOR":"Comparação orientada a dores reais, sem migração forçada",
    }
    return report


@app.get("/api/dashboard")
def dashboard(user=Depends(get_current_user)):
    return dashboard_summary()


@app.get("/api/dashboard/advanced")
def dashboard_full(user=Depends(get_current_user)):
    return dashboard_advanced()


@app.get("/api/followups")
def followups(user=Depends(get_current_user)):
    return load_followups("PENDING", 1000)


@app.post("/api/followups")
def followup_create(payload: FollowupIn, user=Depends(get_current_user)):
    if not get_business(payload.business_key):
        raise HTTPException(404, "Estabelecimento não encontrado")
    fid=create_followup(payload.business_key, payload.title, payload.due_at, "MANUAL", payload.notes)
    audit(_user_email(user), "CREATE_FOLLOWUP", "followup", str(fid), payload.model_dump())
    return {"id": fid}


@app.post("/api/followups/{followup_id}/complete")
def followup_complete(followup_id: int, user=Depends(get_current_user)):
    complete_followup(followup_id)
    return {"ok": True}


@app.get("/api/goals")
def goals(user=Depends(get_current_user)):
    out = []
    for g in load_goals(True, 200):
        g["progress"] = goal_progress(g["metric"], g["start_at"], g["end_at"])
        out.append(g)
    return out


@app.post("/api/goals")
def goal_create(payload: GoalIn, user=Depends(get_current_user)):
    gid=create_goal(payload.metric, payload.target, payload.start_at, payload.end_at, payload.period, payload.label)
    audit(_user_email(user), "CREATE_GOAL", "goal", str(gid), payload.model_dump())
    return {"id": gid}


@app.post("/api/routes/plan")
def route_plan(payload: RouteIn, user=Depends(get_current_user)):
    excluded=set(payload.excluded_business_keys or [])
    filtered_out=[]; leads=[]
    for key in payload.business_keys:
        lead=get_business(key)
        if key in excluded:
            filtered_out.append({"business_key":key,"name":(lead or {}).get("name"),"reason":"Excluído manualmente desta rota"}); continue
        if not lead:
            filtered_out.append({"business_key":key,"name":key,"reason":"Estabelecimento não encontrado"}); continue
        if lead.get("do_not_contact"):
            filtered_out.append({"business_key":key,"name":lead.get("name"),"reason":"Marcado como não contatar"}); continue
        if lead.get("is_large_chain"):
            filtered_out.append({"business_key":key,"name":lead.get("name"),"reason":"Grande rede: baixa aderência para prospecção local"}); continue
        leads.append(enrich_lead(lead))
    end=(payload.end_lat,payload.end_lon) if payload.end_lat is not None and payload.end_lon is not None else None
    result=plan_route_advanced(
        (payload.origin_lat,payload.origin_lon),leads,payload.mode,os.getenv("ORS_API_KEY") or None,GEOAPIFY_KEY or None,
        (datetime.fromisoformat(payload.start_local) if payload.start_local else datetime.now()),payload.available_minutes,payload.visit_minutes,payload.max_stops,payload.return_to_start,
        payload.strategy,end,payload.fixed_business_keys,payload.strategy=="manual",
    )
    result["omitted"] = filtered_out + list(result.get("omitted") or [])
    if payload.persist:
        data=payload.model_dump(); data["end_lat"]=payload.end_lat;data["end_lon"]=payload.end_lon
        result["saved_route"]=create_route_record(data,result,_user_email(user))
    return result

# -----------------------------------------------------------------------------
# Production workspace: Today, public reports, route editor, territories, ICP,
# reanalysis, observability, exports and audit.
# -----------------------------------------------------------------------------

@app.get('/api/today')
def today(user=Depends(get_current_user)):
    return today_workspace(_user_email(user))


@app.post('/api/reports/business/{business_key}/share')
def report_share(business_key: str, payload: PublicReportIn, user=Depends(get_current_user)):
    if not get_business(business_key):
        raise HTTPException(404, 'Estabelecimento não encontrado')
    return create_public_report(business_key, payload.job_id, _user_email(user), payload.expires_days)


@app.get('/public/reports/{token}')
def report_public(token: str):
    report = get_public_report(token)
    if not report:
        raise HTTPException(404, 'Relatório indisponível ou expirado')
    return report


@app.get('/api/public/qr')
def public_qr(data: str = Query(..., min_length=8, max_length=2000)):
    from urllib.parse import urlparse
    parsed=urlparse(data)
    if parsed.scheme not in {'http','https'} or not parsed.netloc:
        raise HTTPException(400,'URL inválida para QR Code')
    # O QR é gerado no próprio backend; o link público não é enviado a terceiros.
    import qrcode
    img=qrcode.make(data)
    buf=io.BytesIO(); img.save(buf,format='PNG')
    return Response(content=buf.getvalue(),media_type='image/png',headers={'Cache-Control':'private, max-age=300'})


@app.post('/api/leads/{business_key}/post-visit-message')
def post_visit_message(business_key: str, payload: MessageIn, user=Depends(get_current_user)):
    lead = get_business(business_key)
    if not lead:
        raise HTTPException(404, 'Estabelecimento não encontrado')
    return {'message': build_post_visit_message(enrich_lead(lead), payload.notes)}


@app.get('/api/leads/{business_key}/evolution')
def lead_evolution(business_key: str, user=Depends(get_current_user)):
    if not get_business(business_key):
        raise HTTPException(404, 'Estabelecimento não encontrado')
    return business_evolution(business_key)


@app.post('/api/leads/{business_key}/reanalysis/schedule')
def reanalysis_schedule(business_key: str, payload: ReanalysisIn, user=Depends(get_current_user)):
    if not get_business(business_key):
        raise HTTPException(404, 'Estabelecimento não encontrado')
    due = schedule_reanalysis(business_key, payload.days)
    audit(_user_email(user), 'SCHEDULE_REANALYSIS', 'business', business_key, {'due_at': due})
    return {'ok': True, 'next_reanalysis_at': due}


@app.post('/api/leads/{business_key}/reanalysis/run')
def reanalysis_run(business_key: str, user=Depends(get_current_user)):
    lead = get_business(business_key)
    if not lead:
        raise HTTPException(404, 'Estabelecimento não encontrado')
    if not lead.get('website'):
        raise HTTPException(400, 'Este lead não possui site para reanalisar')
    try:
        result = analyze_site(lead['website'], use_cache=False)
        lead.update(result)
        lead['last_reanalysis_at'] = datetime.now(timezone.utc).isoformat()
        upsert_businesses([lead])
        recalculate_business_score(business_key)
        mark_reanalyzed(business_key, 60)
        add_activity(business_key, 'REANALYSIS', 'Presença digital reanalisada', 'DONE', _user_email(user))
        audit(_user_email(user), 'RUN_REANALYSIS', 'business', business_key, {'website': lead.get('website')})
        return enrich_lead(get_business(business_key))
    except Exception as exc:
        raise HTTPException(502, f'Não foi possível reanalisar o site: {exc}')


@app.get('/api/routes')
def routes_list(limit: int = 100, user=Depends(get_current_user)):
    return list_routes(limit)


@app.get('/api/routes/{route_id}')
def route_detail(route_id: int, user=Depends(get_current_user)):
    route = get_route(route_id)
    if not route:
        raise HTTPException(404, 'Rota não encontrada')
    return route


@app.patch('/api/routes/{route_id}/stops')
def route_stops_update(route_id: int, payload: RouteStopsIn, user=Depends(get_current_user)):
    if not get_route(route_id):
        raise HTTPException(404, 'Rota não encontrada')
    return update_route_stops(route_id, payload.stops, _user_email(user))


@app.patch('/api/routes/{route_id}/stops/{business_key}')
def route_stop_update(route_id: int, business_key: str, payload: RouteStopStatusIn, user=Depends(get_current_user)):
    if not get_route(route_id):
        raise HTTPException(404, 'Rota não encontrada')
    return update_route_stop_status(route_id, business_key, payload.status, payload.reason, _user_email(user))


@app.get('/api/territories')
def territories(user=Depends(get_current_user)):
    return list_territories(True)


@app.post('/api/territories')
def territory_create(payload: TerritoryIn, user=Depends(get_current_user)):
    try:
        territory_id = create_territory(payload.model_dump(), _user_email(user))
        return {'id': territory_id}
    except Exception as exc:
        raise HTTPException(400, str(exc))


@app.get('/api/icps')
def icps(user=Depends(get_current_user)):
    return list_icps()


@app.post('/api/icps')
def icp_create(payload: ICPIn, user=Depends(get_current_user)):
    try:
        return {'name': upsert_icp(payload.model_dump(), _user_email(user))}
    except ValueError as exc:
        raise HTTPException(400, str(exc))


@app.get('/api/campaigns/financials')
def campaign_financial(user=Depends(get_current_user)):
    return campaign_financials()


@app.get('/api/audit')
def audit_logs(limit: int = 300, user=Depends(get_current_user)):
    _require_role(user, {'ADMIN','MANAGER'})
    return list_audit_logs(limit)


@app.get('/api/export/full')
def export_all(user=Depends(get_current_user)):
    _require_role(user, {'ADMIN','MANAGER'})
    audit(_user_email(user), 'EXPORT_FULL', 'system', None, {})
    return export_full()


@app.get('/api/usage/providers')
def provider_usage(days: int = 30, user=Depends(get_current_user)):
    return provider_usage_summary(days)




@app.get('/api/inbox')
def commercial_inbox(user=Depends(get_current_user)):
    email=_user_email(user); role=_role_for_user(user); now=datetime.now(timezone.utc)
    followups=load_followups('PENDING',500)
    if role=='SELLER':
        followups=[f for f in followups if not f.get('assigned_to') or (f.get('assigned_to') or '').lower()==email]
    due=[f for f in followups if f.get('due_at') and f.get('due_at')<=now.isoformat()]
    upcoming=[f for f in followups if f not in due][:30]
    changes=list_recent_lead_changes(80)
    if role=='SELLER':
        changes=[c for c in changes if not c.get('assigned_to') or (c.get('assigned_to') or '').lower()==email]
    stalled=[enrich_lead(x) for x in list_businesses(1000) if x.get('pipeline_status') in {'INTERESTED','DEMO','PROPOSAL'} and int(x.get('aging_days') or 0)>=7]
    if role=='SELLER': stalled=[x for x in stalled if not x.get('assigned_to') or (x.get('assigned_to') or '').lower()==email]
    return {'due_followups':due[:50],'upcoming':upcoming[:30],'changes':changes[:50],'stalled':stalled[:40],'role':role}

@app.get('/api/intelligence/overview')
def intelligence_dashboard(user=Depends(get_current_user)):
    return intelligence_overview(list_businesses(5000))

@app.post('/api/intelligence/query')
def intelligence_query(payload: NaturalQueryIn, user=Depends(get_current_user)):
    return natural_filter(payload.query, list_businesses(5000))

@app.post('/api/intelligence/copilot')
def intelligence_copilot(payload: CopilotIn, user=Depends(get_current_user)):
    return copilot_answer(payload.question, list_businesses(5000), _user_email(user))

@app.get('/api/leads/{business_key}/changes')
def lead_changes(business_key: str, user=Depends(get_current_user)):
    return load_lead_changes(business_key, 200)

@app.get('/api/leads/{business_key}/briefing')
def lead_briefing(business_key: str, user=Depends(get_current_user)):
    lead=get_business(business_key)
    if not lead: raise HTTPException(404,'Estabelecimento não encontrado')
    intel=compute_lead_intelligence(lead)
    return {
        'lead':enrich_lead(lead),
        'intelligence':intel,
        'playbook':competitor_playbook(lead),
        'duplicates':duplicate_contact_candidates(business_key,lead.get('phone') or lead.get('contact_phone'),lead.get('website'),lead.get('name'),5),
        'last_activities':load_activities(business_key,5),
    }

@app.get('/api/cadences')
def cadences(user=Depends(get_current_user)):
    return list_cadence_templates(False)

@app.post('/api/cadences')
def cadence_create(payload: CadenceIn, user=Depends(get_current_user)):
    _require_role(user, {'ADMIN','MANAGER'})
    out=save_cadence_template(payload.name,payload.description,payload.steps,_user_email(user));audit(_user_email(user),'UPSERT_CADENCE','cadence',str(out.get('id')),payload.model_dump());return out

@app.post('/api/leads/{business_key}/cadence')
def cadence_assign(business_key: str,payload:AssignCadenceIn,user=Depends(get_current_user)):
    if not get_business(business_key):raise HTTPException(404,'Estabelecimento não encontrado')
    try:cid=assign_cadence(business_key,payload.cadence_id,_user_email(user))
    except ValueError as exc:raise HTTPException(400,str(exc))
    audit(_user_email(user),'ASSIGN_CADENCE','business',business_key,{'cadence_id':payload.cadence_id});return {'ok':True,'lead_cadence_id':cid}

@app.get('/api/automations')
def automations(user=Depends(get_current_user)):
    return list_automation_rules(False)

@app.post('/api/automations')
def automation_create(payload:AutomationIn,user=Depends(get_current_user)):
    _require_role(user, {'ADMIN','MANAGER'})
    rid=save_automation_rule(payload.model_dump(),_user_email(user));audit(_user_email(user),'CREATE_AUTOMATION','automation',str(rid),payload.model_dump());return {'id':rid}

@app.patch('/api/automations/{rule_id}')
def automation_toggle(rule_id:int,payload:AutomationToggleIn,user=Depends(get_current_user)):
    _require_role(user, {'ADMIN','MANAGER'})
    set_automation_active(rule_id,payload.active);audit(_user_email(user),'TOGGLE_AUTOMATION','automation',str(rule_id),{'active':payload.active});return {'ok':True}

@app.get('/api/webhooks')
def webhooks(user=Depends(get_current_user)):
    return list_webhooks(False)

@app.post('/api/webhooks')
def webhook_create(payload:WebhookIn,user=Depends(get_current_user)):
    _require_role(user, {'ADMIN','MANAGER'})
    try:url=validate_webhook_url(payload.url)
    except ValueError as exc:raise HTTPException(400,str(exc))
    wid=save_webhook({**payload.model_dump(),'url':url},_user_email(user));audit(_user_email(user),'CREATE_WEBHOOK','webhook',str(wid),{'name':payload.name,'url':url,'events':payload.events});return {'id':wid}

@app.get('/api/webhooks/deliveries')
def webhook_deliveries(limit:int=100,user=Depends(get_current_user)):
    return list_deliveries(limit)

@app.post('/api/import/leads')
def import_leads(payload:ImportLeadsIn,user=Depends(get_current_user)):
    _require_role(user, {'ADMIN','MANAGER'})
    imported=[];duplicates=0;errors=[]
    seen=set()
    for i,raw in enumerate(payload.rows):
        try:
            name=str(raw.get('name') or raw.get('nome') or '').strip()
            if not name:raise ValueError('nome ausente')
            lat=float(raw.get('lat') or raw.get('latitude'))
            lon=float(raw.get('lon') or raw.get('longitude'))
            row={'provider':'Importação','source_id':f'import:{i}','name':name,'address':raw.get('address') or raw.get('endereco'),'lat':lat,'lon':lon,
                 'category':raw.get('category') or raw.get('categoria'),'phone':raw.get('phone') or raw.get('telefone'),'website':raw.get('website') or raw.get('site'),
                 'campaign_name':raw.get('campaign_name') or raw.get('campanha') or 'Importação','estimated_mrr':float(raw.get('estimated_mrr') or raw.get('mrr') or _settings().get('default_mrr') or 89.90),
                 'data_confidence':70,'source_notes':f'Importado de {payload.filename}','raw':raw}
            row['business_key']=compute_identity_key(row);row['identity_key']=compute_identity_key(row);row['source_keys']=[f'import:{payload.filename}:{i}']
            if row['identity_key'] in seen:duplicates+=1;continue
            seen.add(row['identity_key']);imported.append(row)
        except Exception as exc:errors.append({'row':i+1,'error':str(exc)})
    resolve_identity_rows(imported)
    from db import decorate_business_intelligence
    decorate_business_intelligence(imported);upsert_businesses(imported)
    iid=record_import_batch(payload.filename,_user_email(user),len(payload.rows),len(imported),duplicates,len(errors),{'errors':errors[:50]})
    audit(_user_email(user),'IMPORT_LEADS','import',str(iid),{'total':len(payload.rows),'imported':len(imported),'duplicates':duplicates,'errors':len(errors)})
    return {'batch_id':iid,'total':len(payload.rows),'imported':len(imported),'duplicates':duplicates,'errors':errors[:50]}

@app.post('/api/import/file')
def import_file(payload: ImportFileIn, user=Depends(get_current_user)):
    _require_role(user, {'ADMIN','MANAGER'})
    name=payload.filename.lower().strip()
    try:
        raw=base64.b64decode(payload.content_base64,validate=True)
    except Exception:
        raise HTTPException(400,'Arquivo inválido')
    if len(raw)>12*1024*1024:
        raise HTTPException(413,'Arquivo maior que 12 MB')
    rows=[]
    if name.endswith('.xlsx'):
        try:
            from openpyxl import load_workbook
            wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
            ws=wb.active
            values=ws.iter_rows(values_only=True)
            headers=[str(x or '').strip().lower() for x in next(values)]
            for vals in values:
                if len(rows)>=5000: break
                row={headers[i]:vals[i] for i in range(min(len(headers),len(vals))) if headers[i]}
                if any(v not in (None,'') for v in row.values()): rows.append(row)
        except Exception as exc:
            raise HTTPException(400,f'Não foi possível ler o XLSX: {exc}')
    elif name.endswith('.csv'):
        try:
            import csv
            text=raw.decode('utf-8-sig')
            sample=text[:4096]
            dialect=csv.Sniffer().sniff(sample,delimiters=',;\t')
            rows=list(csv.DictReader(io.StringIO(text),dialect=dialect))[:5000]
        except Exception as exc:
            raise HTTPException(400,f'Não foi possível ler o CSV: {exc}')
    else:
        raise HTTPException(400,'Formato suportado: CSV ou XLSX')
    return import_leads(ImportLeadsIn(filename=payload.filename,rows=rows),user)

@app.get('/api/imports')
def imports(user=Depends(get_current_user)):
    return list_import_batches(100)

@app.get('/api/admin/roles')
def roles(user=Depends(get_current_user)):
    rows=list_user_roles()
    if not rows:
        return [{'email':_user_email(user),'role':'ADMIN','implicit':True}]
    return rows

@app.post('/api/admin/roles')
def role_save(payload:RoleIn,user=Depends(get_current_user)):
    email=_user_email(user)
    existing=list_user_roles()
    # O primeiro administrador implícito é persistido antes de criar outros papéis,
    # evitando que ele perca acesso assim que a tabela deixa de estar vazia.
    if not existing and email:
        upsert_user_role(email,'ADMIN')
        existing=list_user_roles()
    _require_role(user, {'ADMIN'})
    target=payload.email.lower().strip()
    if target==email and payload.role!='ADMIN':
        other_admin=any((r.get('email') or '').lower()!=email and r.get('role')=='ADMIN' for r in existing)
        if not other_admin:
            raise HTTPException(400,'Cadastre outro administrador antes de remover seu próprio papel de Admin')
    out=upsert_user_role(target,payload.role);audit(email,'UPSERT_ROLE','user',target,{'role':payload.role});return out

@app.get('/api/admin/overview')
def admin_overview(user=Depends(get_current_user)):
    _require_role(user, {'ADMIN','MANAGER'})
    return {'providers':provider_usage_summary(30),'ai':ai_usage_today(),'audits':list_audit_logs(50),'imports':list_import_batches(30),'roles':list_user_roles(),'privacy':list_privacy_requests(50),'webhook_deliveries':list_deliveries(50)}

@app.get('/api/privacy/requests')
def privacy_list(user=Depends(get_current_user)):
    _require_role(user, {'ADMIN','MANAGER'})
    return list_privacy_requests(100)

@app.post('/api/privacy/requests')
def privacy_create(payload:PrivacyIn,user=Depends(get_current_user)):
    rid=create_privacy_request(payload.model_dump(),_user_email(user));audit(_user_email(user),'PRIVACY_REQUEST','privacy',str(rid),payload.model_dump());return {'id':rid}


@app.post('/api/privacy/requests/{request_id}/complete')
def privacy_complete(request_id:int,payload:PrivacyCompleteIn,user=Depends(get_current_user)):
    _require_role(user, {'ADMIN','MANAGER'})
    try:out=complete_privacy_request(request_id,_user_email(user),payload.anonymize)
    except ValueError as exc:raise HTTPException(404,str(exc))
    audit(_user_email(user),'COMPLETE_PRIVACY_REQUEST','privacy',str(request_id),{'anonymize':payload.anonymize})
    return out

@app.get('/api/health/providers')
def health_providers(user=Depends(get_current_user)):
    jobs = list_search_jobs(10)
    return {
        'database': 'configured' if os.getenv('DATABASE_URL') else 'local',
        'geoapify': {'configured': bool(GEOAPIFY_KEY)},
        'foursquare': {'configured': bool(FOURSQUARE_KEY)},
        'ai': {k: {'configured': bool(v), 'model': _ai_runtime()[0].get(k)} for k, v in _AI_KEYS.items()},
        'routing': {'geoapify': bool(GEOAPIFY_KEY), 'openrouteservice': bool(os.getenv('ORS_API_KEY')), 'fallback': 'OSRM/local'},
        'jobs': {
            'active': sum(1 for j in jobs if j.get('status') in {'RUNNING','QUEUED'}),
            'failed_recent': sum(1 for j in jobs if j.get('status') == 'FAILED'),
        },
        'ai_usage_today': ai_usage_today(),
    }
